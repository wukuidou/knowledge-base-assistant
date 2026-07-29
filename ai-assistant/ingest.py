"""文档导入器 v2.0 — 支持 PDF / Word / HTML / 代码 / 纯文本"""
import os
os.environ["HF_ENDPOINT"] = "https://hf-mirror.com"

import fitz  # pymupdf — PDF
from pathlib import Path
from dotenv import load_dotenv
from llama_index.core import (
    VectorStoreIndex,
    Settings,
    StorageContext,
    Document,
)
from llama_index.core.node_parser import SentenceSplitter
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
from llama_index.vector_stores.chroma import ChromaVectorStore
import chromadb

# ---- 新增格式支持 ----
from docx import Document as DocxDocument  # Word
from bs4 import BeautifulSoup             # HTML

load_dotenv()

# 1. 设置 Embedding 模型
Settings.embed_model = HuggingFaceEmbedding(
    model_name="BAAI/bge-small-zh-v1.5",
)

# 2. 初始化 ChromaDB
db = chromadb.PersistentClient(path="./storage/chroma_db")
chroma_collection = db.get_or_create_collection("knowledge_base")
vector_store = ChromaVectorStore(chroma_collection=chroma_collection)
storage_context = StorageContext.from_defaults(vector_store=vector_store)

# 3. 格式路由 — 按扩展名分发到对应解析器
CODE_EXTS = {
    ".py", ".js", ".ts", ".jsx", ".tsx",
    ".java", ".cpp", ".c", ".h", ".hpp",
    ".go", ".rs", ".rb", ".php", ".swift",
    ".kt", ".scala", ".sql", ".sh", ".bash", ".m", ".r",
}


def load_file(filepath: Path) -> list[Document]:
    ext = filepath.suffix.lower()
    if ext == ".pdf":
        return load_pdf(filepath)
    elif ext == ".docx":
        return load_docx(filepath)
    elif ext in (".html", ".htm"):
        return load_html(filepath)
    elif ext in CODE_EXTS:
        return load_code(filepath)
    else:
        return load_text(filepath)


def load_pdf(filepath: Path) -> list[Document]:
    """PyMuPDF 解析 PDF（中文不乱码）"""
    doc = fitz.open(str(filepath))
    text = ""
    for page in doc:
        text += page.get_text()
    doc.close()
    if text.strip():
        return [Document(text=text, metadata={
            "file_name": filepath.name, "source": str(filepath), "type": "pdf",
        })]
    return []


def load_docx(filepath: Path) -> list[Document]:
    """python-docx 解析 Word 文档"""
    doc = DocxDocument(str(filepath))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    # 也提取表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            text += "\n" + " | ".join(cells)
    if text.strip():
        return [Document(text=text, metadata={
            "file_name": filepath.name, "source": str(filepath), "type": "docx",
        })]
    return []


def load_html(filepath: Path) -> list[Document]:
    """BeautifulSoup 解析 HTML 提取纯文本"""
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        soup = BeautifulSoup(f.read(), "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    text = "\n\n".join(lines)
    if text.strip():
        return [Document(text=text, metadata={
            "file_name": filepath.name, "source": str(filepath), "type": "html",
        })]
    return []


def load_code(filepath: Path) -> list[Document]:
    """加载代码文件，用更小的 chunk"""
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    if text.strip():
        return [Document(text=text, metadata={
            "file_name": filepath.name,
            "source": str(filepath),
            "type": "code",
            "language": filepath.suffix[1:],
        })]
    return []


def load_text(filepath: Path) -> list[Document]:
    """加载纯文本 / 无扩展名文件"""
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    if text.strip():
        return [Document(text=text, metadata={
            "file_name": filepath.name, "source": str(filepath), "type": "text",
        })]
    return []


# 4. 扫描并加载所有文件
data_dir = Path("data")
all_documents = []
file_stats: dict[str, int] = {}

for filepath in sorted(data_dir.rglob("*")):
    if not filepath.is_file():
        continue
    # 跳过隐藏文件和临时文件
    if filepath.name.startswith(".") or filepath.suffix.lower() in (".tmp", ".temp"):
        continue
    try:
        docs = load_file(filepath)
        all_documents.extend(docs)
        ext = filepath.suffix.lower() or "(无后缀)"
        file_stats[ext] = file_stats.get(ext, 0) + 1
        if docs:
            print(f"  [OK] {filepath.suffix or '  ':6s} {filepath.name}")
        else:
            print(f"  [EMPTY] {filepath.name}")
    except Exception as e:
        print(f"  [ERR] {filepath.name} — {e}")

total = len(all_documents)
print(f"\n[DOCS] 共加载 {total} 个文档")
if file_stats:
    print("   格式分布:", ", ".join(f"{ext}: {n}" for ext, n in sorted(file_stats.items())))

if total == 0:
    print("[WARN] 没有需要处理的文档，退出。")
    exit(0)

# 5. 分块：代码文件用更小的 chunk_size=256，其他用 512
code_docs = [d for d in all_documents if d.metadata.get("type") == "code"]
other_docs = [d for d in all_documents if d.metadata.get("type") != "code"]

all_nodes = []

if code_docs:
    code_parser = SentenceSplitter(chunk_size=256, chunk_overlap=30)
    code_nodes = code_parser.get_nodes_from_documents(code_docs)
    print(f"  [CODE] 代码文件: {len(code_docs)} 个 -> {len(code_nodes)} 个块 (chunk_size=256)")
    all_nodes.extend(code_nodes)

if other_docs:
    text_parser = SentenceSplitter(chunk_size=512, chunk_overlap=50)
    other_nodes = text_parser.get_nodes_from_documents(other_docs)
    print(f"  [TEXT] 其他文件: {len(other_docs)} 个 -> {len(other_nodes)} 个块 (chunk_size=512)")
    all_nodes.extend(other_nodes)

# 6. 构建索引
index = VectorStoreIndex(
    nodes=all_nodes,
    storage_context=storage_context,
    show_progress=True,
)

count = chroma_collection.count()
print(f"\n[DONE] 索引构建完成！")
print(f"[STATS] 向量库中文档块数: {count}")
